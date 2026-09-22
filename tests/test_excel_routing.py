# -*- coding: utf-8 -*-
from types import SimpleNamespace

from openpyxl import Workbook

import src.parsers.smart_parser as smart_parser
from src.parsers.smart_parser import ParsedDocument, parse_xlsx


def _fake_config(max_mb):
    return SimpleNamespace(
        document=SimpleNamespace(excel_mineru_max_mb=max_mb)
    )


def test_parse_file_routes_large_xlsx_to_openpyxl(monkeypatch, tmp_path):
    large_path = tmp_path / "large.xlsx"
    large_path.write_bytes(b"x" * 1024)
    captured = {}

    monkeypatch.setattr(
        smart_parser,
        "get_config",
        lambda: _fake_config(0.000001),
    )

    def fake_parse_xlsx(path, prefer_openpyxl):
        captured["path"] = path
        captured["prefer_openpyxl"] = prefer_openpyxl
        return ParsedDocument(
            doc_id="xlsx:large",
            source="large.xlsx",
            file_type="xlsx",
            extra={"routed": True},
        )

    monkeypatch.setattr(
        smart_parser,
        "parse_xlsx",
        fake_parse_xlsx,
    )

    def fail_mineru(path):
        raise AssertionError("MinerU should not be used")

    monkeypatch.setattr(
        smart_parser,
        "parse_with_mineru",
        fail_mineru,
    )

    parsed = smart_parser.parse_file(str(large_path))

    assert parsed.doc_id == "xlsx:large"
    assert parsed.extra["routed"] is True
    assert captured["prefer_openpyxl"] is True


def test_parse_file_keeps_small_xlsx_on_mineru(monkeypatch, tmp_path):
    small_path = tmp_path / "small.xlsx"
    small_path.write_bytes(b"x" * 16)
    captured = {}

    monkeypatch.setattr(smart_parser, "get_config", lambda: _fake_config(100))

    def fake_mineru(path):
        captured["path"] = path
        return ParsedDocument(
            doc_id="mineru:small",
            source="small.xlsx",
            file_type="xlsx",
        )

    monkeypatch.setitem(
        smart_parser._PARSER_REGISTRY,
        ".xlsx",
        fake_mineru,
    )

    def fail_xlsx(*args, **kwargs):
        raise AssertionError("openpyxl should not be used")

    monkeypatch.setattr(
        smart_parser,
        "parse_xlsx",
        fail_xlsx,
    )

    parsed = smart_parser.parse_file(str(small_path))

    assert parsed.doc_id == "mineru:small"
    assert captured["path"] == str(small_path)


def test_parse_xlsx_prefer_openpyxl_streams_rows_by_sheet(tmp_path):
    excel_path = tmp_path / "inventory.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "库存"
    ws.append(["项目", "数量"])
    ws.append(["白酒", 100])
    wb.save(excel_path)

    parsed = parse_xlsx(str(excel_path), prefer_openpyxl=True)

    assert parsed.file_type == "xlsx"
    assert parsed.extra["parser_backend"] == "openpyxl"
    assert len(parsed.blocks) == 1
    assert parsed.blocks[0].block_type == "table"
    assert parsed.blocks[0].content == "项目: 白酒 | 数量: 100"
    assert parsed.blocks[0].extra == {"sheet": "库存", "row": 2}
